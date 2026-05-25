import json

import pytest

from services.parsing.excel_parser import parse_excel
from services.parsing.router import parse_document


@pytest.fixture
def sample_xlsx_path(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "requirements.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Payments"
    ws1.append(["Requirement ID", "Description", "Priority"])
    ws1.append(["REQ-001", "Process invoices", "High"])
    ws1.append(["REQ-002", "Send reminders", "Medium"])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Total", "2"])
    wb.save(path)
    wb.close()
    return path


def test_parse_excel_two_sheets_and_row_count(sample_xlsx_path):
    elements = parse_excel(str(sample_xlsx_path))
    sheets = [e for e in elements if e.category == "SheetTitle"]
    rows = [e for e in elements if e.category == "TableRow"]
    assert len(sheets) == 2
    assert {s.sheet_name for s in sheets} == {"Payments", "Summary"}
    assert len(rows) == 3


def test_table_row_json_preserves_columns(sample_xlsx_path):
    elements = parse_excel(str(sample_xlsx_path))
    req_rows = [
        e for e in elements if e.category == "TableRow" and e.sheet_name == "Payments"
    ]
    assert len(req_rows) == 2
    payload = json.loads(req_rows[0].text)
    assert payload["sheet"] == "Payments"
    assert payload["Requirement ID"] == "REQ-001"
    assert payload["Description"] == "Process invoices"
    assert payload["Priority"] == "High"
    assert req_rows[0].row_index == 2


def test_parse_excel_markdown_format(sample_xlsx_path):
    elements = parse_excel(str(sample_xlsx_path), row_format="markdown")
    row = next(e for e in elements if e.category == "TableRow" and e.row_index == 2)
    assert "| Requirement ID | Description | Priority |" in row.text
    assert "| REQ-001 | Process invoices | High |" in row.text


def test_parse_document_routes_xlsx_to_excel_parser(sample_xlsx_path):
    elements = parse_document(str(sample_xlsx_path), "xlsx")
    assert any(e.category == "TableRow" for e in elements)
    assert any(e.category == "SheetTitle" for e in elements)


def test_router_does_not_select_excel_for_pdf():
    from services.parsing.excel_parser import EXCEL_FILE_TYPES
    from services.parsing.router import normalize_file_type

    assert normalize_file_type("pdf") not in EXCEL_FILE_TYPES
    assert normalize_file_type("xlsx") in EXCEL_FILE_TYPES


def test_empty_sheet_skipped(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "mixed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["A", "B"])
    ws.append([1, 2])
    empty = wb.create_sheet("Empty")
    wb.save(path)
    wb.close()
    elements = parse_excel(str(path))
    assert not any(e.sheet_name == "Empty" for e in elements)


def test_duplicate_header_names_normalized(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "dup.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Name", "Value"])
    ws.append(["a", "b", "1"])
    wb.save(path)
    wb.close()
    row = parse_excel(str(path))[-1]
    data = json.loads(row.text)
    assert "Name" in data
    assert "Name_2" in data

"""Parse Excel workbooks into sheet-aware, row-level elements for table chunking."""

from pathlib import Path

from openpyxl import load_workbook

from .table_rows import cell_str, row_is_empty, rows_to_sheet_elements
from .types import ParsedElement

EXCEL_FILE_TYPES = frozenset({"xlsx", "xlsm"})
TABLE_FILE_TYPES = frozenset({"xlsx", "xlsm", "xls", "csv"})


def parse_excel(
    file_path: str,
    *,
    row_format: str = "json",
) -> list[ParsedElement]:
    """
    Parse .xlsx/.xlsm into ordered elements:
    SheetTitle -> TableHeader (markdown header row) -> TableRow per data row.
    """
    path = Path(file_path)
    if path.suffix.lower().replace(".", "") not in EXCEL_FILE_TYPES:
        raise ValueError(f"Unsupported Excel type: {path.suffix}")

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    elements: list[ParsedElement] = []

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_raw: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows_raw.append([cell_str(c) for c in row])

            elements.extend(
                rows_to_sheet_elements(sheet_name, rows_raw, row_format=row_format)
            )
    finally:
        wb.close()

    return elements
